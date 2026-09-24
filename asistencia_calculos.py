"""Cálculo de horas con recargo 50%, 25% y 100% desde eventos de asistencia."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from typing import Dict, Iterable, List, Sequence, Tuple

from asistencia_pdf import EventoAsistencia

# Horas extras diurnas después de las 17:00 (50%)
HORA_LIMITE_50 = time(17, 0)

# Franja nocturna para recargo 25% (turno / permanece de noche)
HORA_NOCHE_INICIO = time(19, 0)
HORA_NOCHE_FIN = time(6, 0)

# Duración plausible de un turno inferido entre dos marcas consecutivas
MIN_SESSION_HORAS = 4.0
MAX_SESSION_HORAS = 13.0

# Feriados configurables (Agosto 2026 — Ecuador: 10 de agosto)
FERIADOS_DEFAULT: Tuple[date, ...] = (
    date(2026, 8, 10),
)


@dataclass(frozen=True)
class SesionTrabajo:
    id_empleado: str
    nombre: str
    inicio: datetime
    fin: datetime
    horas_totales: Decimal
    horas_50: Decimal
    horas_25: Decimal
    horas_100: Decimal
    es_turno_noche: bool


@dataclass(frozen=True)
class ResumenOperario:
    id_empleado: str
    nombre: str
    horas_totales: Decimal
    horas_50: Decimal
    horas_25: Decimal
    horas_100: Decimal
    sesiones: int


def _parse_fecha_hora(texto: str) -> datetime:
    return datetime.strptime(texto.strip(), "%d/%m/%Y %H:%M")


def _decimal_horas(segundos: float) -> Decimal:
    return (Decimal(str(segundos)) / Decimal("3600")).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )


def _es_feriado_o_fin_de_semana(d: date, feriados: Sequence[date]) -> bool:
    return d.weekday() >= 5 or d in feriados


def _es_turno_noche(inicio: datetime, fin: datetime) -> bool:
    """Turno que cruza medianoche o sale en madrugada tras entrada vespertina."""
    if fin.date() > inicio.date():
        return True
    if inicio.time() >= HORA_LIMITE_50 and fin.time() <= time(10, 0):
        return True
    return False


def _en_franja_noche(dt: datetime) -> bool:
    t = dt.time()
    return t >= HORA_NOCHE_INICIO or t < HORA_NOCHE_FIN


def inferir_sesiones(
    eventos: Sequence[EventoAsistencia],
) -> List[Tuple[EventoAsistencia, datetime, datetime]]:
    """
    Infiere sesiones entre marcas consecutivas con duración de turno plausible.
    """
    por_persona: Dict[Tuple[str, str], List[EventoAsistencia]] = {}
    for ev in eventos:
        clave = (ev.id_empleado, ev.nombre)
        por_persona.setdefault(clave, []).append(ev)

    sesiones: List[Tuple[EventoAsistencia, datetime, datetime]] = []
    for (id_emp, nombre), lista in por_persona.items():
        lista_ord = sorted(lista, key=lambda e: _parse_fecha_hora(e.fecha_hora))
        for i in range(len(lista_ord) - 1):
            ev = lista_ord[i]
            t1 = _parse_fecha_hora(ev.fecha_hora)
            t2 = _parse_fecha_hora(lista_ord[i + 1].fecha_hora)
            horas = (t2 - t1).total_seconds() / 3600
            if MIN_SESSION_HORAS <= horas <= MAX_SESSION_HORAS:
                sesiones.append((ev, t1, t2))
    return sesiones


def clasificar_sesion(
    id_empleado: str,
    nombre: str,
    inicio: datetime,
    fin: datetime,
    feriados: Sequence[date],
) -> SesionTrabajo:
    """
    Clasifica horas de una sesión:
    - 100%: fines de semana y feriados (prioridad)
    - 25%: turno nocturno / permanece de noche (días laborables)
    - 50%: horas extras después de las 17:00 (días laborables, fuera de turno noche)
    """
    turno_noche = _es_turno_noche(inicio, fin)
    h50 = Decimal("0")
    h25 = Decimal("0")
    h100 = Decimal("0")

    cursor = inicio
    paso = timedelta(minutes=1)
    while cursor < fin:
        fin_tramo = min(cursor + paso, fin)
        seg = (fin_tramo - cursor).total_seconds()
        fraccion = Decimal(str(seg)) / Decimal("3600")
        d = cursor.date()

        if _es_feriado_o_fin_de_semana(d, feriados):
            h100 += fraccion
        elif turno_noche and _en_franja_noche(cursor):
            h25 += fraccion
        elif cursor.time() >= HORA_LIMITE_50:
            h50 += fraccion
        cursor = fin_tramo

    total = _decimal_horas((fin - inicio).total_seconds())
    return SesionTrabajo(
        id_empleado=id_empleado,
        nombre=nombre,
        inicio=inicio,
        fin=fin,
        horas_totales=total,
        horas_50=h50.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        horas_25=h25.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        horas_100=h100.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        es_turno_noche=turno_noche,
    )


def calcular_recargos(
    eventos: Sequence[EventoAsistencia],
    feriados: Sequence[date] | None = None,
) -> Tuple[List[SesionTrabajo], List[ResumenOperario]]:
    feriados = feriados or FERIADOS_DEFAULT
    sesiones_raw = inferir_sesiones(eventos)
    sesiones: List[SesionTrabajo] = []
    for ev, t1, t2 in sesiones_raw:
        sesiones.append(clasificar_sesion(ev.id_empleado, ev.nombre, t1, t2, feriados))

    resumen_map: Dict[Tuple[str, str], ResumenOperario] = {}
    for s in sesiones:
        clave = (s.id_empleado, s.nombre)
        if clave not in resumen_map:
            resumen_map[clave] = ResumenOperario(
                id_empleado=s.id_empleado,
                nombre=s.nombre,
                horas_totales=Decimal("0"),
                horas_50=Decimal("0"),
                horas_25=Decimal("0"),
                horas_100=Decimal("0"),
                sesiones=0,
            )
        r = resumen_map[clave]
        resumen_map[clave] = ResumenOperario(
            id_empleado=r.id_empleado,
            nombre=r.nombre,
            horas_totales=(r.horas_totales + s.horas_totales).quantize(Decimal("0.01")),
            horas_50=(r.horas_50 + s.horas_50).quantize(Decimal("0.01")),
            horas_25=(r.horas_25 + s.horas_25).quantize(Decimal("0.01")),
            horas_100=(r.horas_100 + s.horas_100).quantize(Decimal("0.01")),
            sesiones=r.sesiones + 1,
        )

    resumen = sorted(resumen_map.values(), key=lambda x: x.nombre.lower())
    return sesiones, resumen


def agregar_hojas_recargo_excel(
    ruta_excel: str,
    sesiones: Sequence[SesionTrabajo],
    resumen: Sequence[ResumenOperario],
) -> None:
    """Añade hojas Sesiones y Resumen recargos a un libro existente."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(ruta_excel)

    if "Sesiones" in wb.sheetnames:
        del wb["Sesiones"]
    if "Resumen recargos" in wb.sheetnames:
        del wb["Resumen recargos"]

    ws_s = wb.create_sheet("Sesiones")
    headers_s = [
        "ID",
        "Nombre",
        "Inicio",
        "Fin",
        "Horas total",
        "Horas 50%",
        "Horas 25%",
        "Horas 100%",
        "Turno noche",
    ]
    bold = Font(bold=True)
    for c, h in enumerate(headers_s, 1):
        ws_s.cell(1, c, h).font = bold
    for i, s in enumerate(sesiones, 2):
        ws_s.cell(i, 1, s.id_empleado)
        ws_s.cell(i, 2, s.nombre)
        ws_s.cell(i, 3, s.inicio.strftime("%d/%m/%Y %H:%M"))
        ws_s.cell(i, 4, s.fin.strftime("%d/%m/%Y %H:%M"))
        ws_s.cell(i, 5, float(s.horas_totales))
        ws_s.cell(i, 6, float(s.horas_50))
        ws_s.cell(i, 7, float(s.horas_25))
        ws_s.cell(i, 8, float(s.horas_100))
        ws_s.cell(i, 9, "Sí" if s.es_turno_noche else "No")

    ws_r = wb.create_sheet("Resumen recargos")
    headers_r = [
        "ID",
        "Nombre",
        "Sesiones",
        "Horas total",
        "Horas 50%",
        "Horas 25%",
        "Horas 100%",
    ]
    for c, h in enumerate(headers_r, 1):
        ws_r.cell(1, c, h).font = bold
    for i, r in enumerate(resumen, 2):
        ws_r.cell(i, 1, r.id_empleado)
        ws_r.cell(i, 2, r.nombre)
        ws_r.cell(i, 3, r.sesiones)
        ws_r.cell(i, 4, float(r.horas_totales))
        ws_r.cell(i, 5, float(r.horas_50))
        ws_r.cell(i, 6, float(r.horas_25))
        ws_r.cell(i, 7, float(r.horas_100))

    wb.save(ruta_excel)


def procesar_pdf_completo(
    ruta_pdf: str,
    ruta_excel: str | None = None,
    feriados: Sequence[date] | None = None,
) -> Tuple[str, List[ResumenOperario]]:
    """PDF → Excel (Eventos + Sesiones + Resumen recargos)."""
    from asistencia_pdf import extraer_eventos_desde_pdf, exportar_eventos_a_excel

    eventos = extraer_eventos_desde_pdf(ruta_pdf)
    if ruta_excel is None:
        from pathlib import Path

        ruta_excel = str(Path(ruta_pdf).with_suffix(".xlsx"))
    exportar_eventos_a_excel(eventos, ruta_excel)
    sesiones, resumen = calcular_recargos(eventos, feriados)
    agregar_hojas_recargo_excel(ruta_excel, sesiones, resumen)
    return ruta_excel, list(resumen)
