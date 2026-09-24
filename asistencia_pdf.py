"""Extracción de reportes de asistencia PDF a Excel."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Sequence

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font

COLUMNAS = ("ID", "Nombre", "Fecha / Hora", "Estado", "Tipo de Registro")


@dataclass(frozen=True)
class EventoAsistencia:
    id_empleado: str
    nombre: str
    fecha_hora: str
    estado: str
    tipo: str


def extraer_eventos_desde_pdf(ruta_pdf: str | Path) -> List[EventoAsistencia]:
    """Lee todas las páginas del PDF y devuelve la lista de eventos."""
    ruta = Path(ruta_pdf)
    if not ruta.is_file():
        raise FileNotFoundError(f"No se encontró el PDF: {ruta}")

    eventos: List[EventoAsistencia] = []
    with pdfplumber.open(ruta) as pdf:
        for pagina in pdf.pages:
            for tabla in pagina.extract_tables() or []:
                for fila in tabla:
                    if not fila or not fila[0] or fila[0] == "ID":
                        continue
                    if not fila[2]:
                        continue
                    eventos.append(
                        EventoAsistencia(
                            id_empleado=str(fila[0]).strip(),
                            nombre=(fila[1] or "").strip(),
                            fecha_hora=str(fila[2]).strip(),
                            estado=(fila[3] or "").strip(),
                            tipo=(fila[4] or "Normal").strip(),
                        )
                    )
    return eventos


def exportar_eventos_a_excel(
    eventos: Sequence[EventoAsistencia],
    ruta_salida: str | Path,
) -> Path:
    """Exporta eventos al mismo formato de columnas que el PDF."""
    salida = Path(ruta_salida)
    wb = Workbook()
    ws = wb.active
    ws.title = "Eventos"

    header_font = Font(bold=True)
    for col, titulo in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.font = header_font

    for i, ev in enumerate(eventos, start=2):
        ws.cell(row=i, column=1, value=ev.id_empleado)
        ws.cell(row=i, column=2, value=ev.nombre)
        ws.cell(row=i, column=3, value=ev.fecha_hora)
        ws.cell(row=i, column=4, value=ev.estado)
        ws.cell(row=i, column=5, value=ev.tipo)

    from openpyxl.utils import get_column_letter

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 22

    wb.save(salida)
    return salida


def pdf_a_excel(ruta_pdf: str | Path, ruta_excel: str | Path | None = None) -> Path:
    """Convierte un PDF de asistencia a Excel (hoja Eventos)."""
    pdf_path = Path(ruta_pdf)
    if ruta_excel is None:
        ruta_excel = pdf_path.with_suffix(".xlsx")
    eventos = extraer_eventos_desde_pdf(pdf_path)
    return exportar_eventos_a_excel(eventos, ruta_excel)
